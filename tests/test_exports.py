import io
import pytest
from docx import Document
from document_utils.txt_generator import format_txt
from document_utils.docx_generator import format_docx
from document_utils.pdf_generator import format_pdf

SAMPLE_DOC = """## Employment Contract

This Employment Agreement is entered into on April 15, 2025 by and between:

Jane Doe (Employee) and Acme Global Corp. (Employer)

### 1. Position and Duties
The Employee will serve as Senior Systems Architect.

### 2. Compensation & Terms
1. Base salary of $120,000 paid bi-weekly.
2. Standard health insurance and 20 days paid annual leave.
3. 30 days notice required for termination.

### 3. Confidentiality
Employee will maintain complete confidentiality of trade secrets.

IN WITNESS WHEREOF, the parties have signed below.

______________________________________            ______________________________________
Authorized Signature                              Authorized Signature
"""


def test_txt_export():
    txt_bytes = format_txt(SAMPLE_DOC, "Employment Contract")
    assert isinstance(txt_bytes, bytes)
    assert len(txt_bytes) > 50
    content = txt_bytes.decode("utf-8")
    assert "EMPLOYMENT CONTRACT" in content
    assert "Senior Systems Architect" in content
    assert "LegalEase Draft" in content


def test_docx_export():
    docx_bytes = format_docx(SAMPLE_DOC, "Employment Contract", org_name="Test Legal Org")
    assert isinstance(docx_bytes, bytes)
    assert len(docx_bytes) > 500

    # Parse DOCX in-memory to verify structure
    stream = io.BytesIO(docx_bytes)
    doc = Document(stream)
    
    # Verify title paragraph
    paragraphs = [p.text for p in doc.paragraphs if p.text]
    assert any("EMPLOYMENT CONTRACT" in p for p in paragraphs)
    assert any("Senior Systems Architect" in p for p in paragraphs)

    # Verify signature table was created
    assert len(doc.tables) >= 1
    table = doc.tables[0]
    table_text = " ".join(cell.text for row in table.rows for cell in row.cells)
    assert "Authorized Signature" in table_text


def test_pdf_export():
    pdf_bytes = format_pdf(SAMPLE_DOC, "Employment Contract", org_name="Test Legal Org")
    assert isinstance(pdf_bytes, bytes)
    assert len(pdf_bytes) > 500
    # PDF files start with the magic header '%PDF-'
    assert pdf_bytes.startswith(b"%PDF-")


def test_pdf_multipage_handling():
    # Generate long legal document text that spans multiple pages
    long_doc = SAMPLE_DOC + "\n\n" + "\n\n".join(
        f"### {i}. Additional Clause {i}\n"
        f"This is an extensive legal covenant ensuring full compliance with subsection {i}."
        f"The parties agree that all representations in clause {i} remain binding."
        for i in range(4, 30)
    )
    pdf_bytes = format_pdf(long_doc, "Long Employment Contract")
    assert isinstance(pdf_bytes, bytes)
    assert pdf_bytes.startswith(b"%PDF-")
    # Multi-page PDF should be substantially larger
    assert len(pdf_bytes) > 5000
